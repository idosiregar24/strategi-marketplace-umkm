"""Ekspor data_clean dan data_features ke satu berkas Excel untuk dibagikan: `python -m umkm.ekspor`.

Isi: sheet Petunjuk, Data bersih, Data siap analisis, dan Kamus data. Kolom `nama_usaha` tidak ikut
(identitas usaha); responden dikenali lewat `id_responden` (R001-R149, urut sesuai baris Excel asli).
Berkas ditulis ke data/processed/ (tidak masuk repo).
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from umkm import config

KOLOM_PRIBADI = ["nama_usaha"]
NAMA_BERKAS = "data_umkm.xlsx"

BUTIR = {
    "x1_unggah": (
        "X1 Unggah produk",
        [
            "unggah rutin mingguan",
            "unggah produk baru",
            "unggah ulang produk",
            "pantau produk unggahan",
        ],
    ),
    "x2_promosi": (
        "X2 Promosi",
        ["voucher/diskon", "iklan berbayar", "fitur promosi", "evaluasi promosi"],
    ),
    "x3_komunikasi": (
        "X3 Komunikasi pelanggan",
        [
            "respon cepat",
            "jelas dan sopan",
            "balas di hari yang sama",
            "tangani komplain",
            "jaga komunikasi",
        ],
    ),
    "x4_konsistensi": (
        "X4 Konsistensi toko",
        [
            "toko selalu aktif",
            "stok diperbarui",
            "foto/deskripsi terjaga",
            "harga akurat",
            "tetap kelola saat sibuk",
        ],
    ),
    "y_persepsi": (
        "Persepsi kinerja (pendukung, bukan target)",
        [
            "omzet naik",
            "penjualan tumbuh",
            "pesanan naik",
            "transaksi lebih sering",
            "pelanggan baru",
            "jangkauan luar daerah",
            "usaha berkembang",
            "kontribusi marketplace",
        ],
    ),
}
PLATFORM = {
    "gofood": "GoFood",
    "grabfood": "GrabFood",
    "shopee": "Shopee",
    "maxim": "Maxim",
    "tiktok": "TikTok Shop",
}

KAMUS: dict[str, str] = {
    "id_responden": "Kode responden R001-R149 (urut sesuai baris di Excel kuesioner asli)",
    "timestamp": "Waktu pengisian kuesioner",
    "lama_usaha_pra_mp": "Lama usaha sebelum memakai marketplace (tahun); 0 = berdiri bersamaan",
    "lama_usaha_pra_mp_raw": "Isian asli lama usaha sebelum marketplace",
    "lama_di_mp": "Lama berjualan di marketplace (tahun)",
    "lama_di_mp_raw": "Isian asli lama berjualan di marketplace",
    "platform_raw": "Isian asli platform marketplace yang dipakai",
    "jumlah_platform": "Jumlah platform marketplace yang dipakai",
    "pelanggan_baru": "Pelanggan baru per bulan (isian asli; waktu pengukuran ambigu)",
    "flag_kalender_lama_usaha_pra_mp": "1 = lama usaha diisi tahun kalender (mis. 2018) lalu dikonversi",
    "flag_kalender_lama_di_mp": "1 = lama di marketplace diisi tahun kalender lalu dikonversi",
    "flag_lama_tahun_kalender": "1 = salah satu isian lama usaha berupa tahun kalender",
    "flag_outlier": "1 = ada nilai omzet/transaksi yang dipangkas di persentil 99",
    "flag_trx_dikosongkan": "1 = transaksi janggal sehingga dikosongkan (4 responden)",
    "flag_trx_pecahan": "1 = transaksi bukan bilangan bulat",
    "flag_trx_tidak_wajar": "1 = transaksi > 100.000 atau sama persis dengan omzet",
    "tanpa_baseline": "1 = usaha berdiri bersamaan dengan masuk marketplace (tidak punya kondisi sebelum)",
    "sampel_utama": "1 = dipakai di analisis utama (punya kondisi sebelum marketplace); 116 responden",
    "log_omzet_sebelum": "log(1 + omzet sebelum terpangkas)",
    "skala_omzet": "Kategori skala usaha dari kuartil omzet sebelum (Q1 <= 9 jt, Q2 <= 18 jt, Q3 <= 30 jt, Q4 > 30 jt)",
    "d_omzet": "Perubahan omzet: log(1 + sesudah) - log(1 + sebelum); > 0 berarti naik",
    "d_trx": "Perubahan transaksi: log(1 + sesudah) - log(1 + sebelum); > 0 berarti naik",
    "arah_omzet": "Arah perubahan omzet: naik / sama / turun",
    "arah_trx": "Arah perubahan transaksi: naik / sama / turun / kosong",
    "pct_omzet": "Pertumbuhan omzet (%); kosong bila omzet sebelum = 0",
    "pct_trx": "Pertumbuhan transaksi (%); kosong bila transaksi sebelum = 0 atau kosong",
    "naik_penjualan": "TARGET UTAMA: 1 = omzet ATAU transaksi naik",
    "y_obj": "TARGET: skor gabungan besarnya perubahan omzet dan transaksi (z-score)",
    "y_pelanggan_baru": "TARGET TAMBAHAN: log(1 + pelanggan baru per bulan)",
    "flag_jawaban_seragam": "1 = 18 butir strategi dijawab dengan angka yang sama persis",
    "flag_omzet_turun_tajam": "1 = omzet sesudah <= 20% omzet sebelum (mis. turun tepat 10 kali)",
}


def keterangan(kolom: str) -> str:
    """Penjelasan satu kolom dalam bahasa sederhana."""
    if kolom in KAMUS:
        return KAMUS[kolom]
    if m := re.fullmatch(r"(x[1-4]_\w+?|y_persepsi)_(\d)", kolom):
        nama, butir = BUTIR[m.group(1)]
        i = int(m.group(2))
        return f"{nama}, butir {i}: {butir[i - 1]} (skala 1-5)"
    if kolom in BUTIR:
        return f"Skor {BUTIR[kolom][0]} = rata-rata butir (skala 1-5)"
    if m := re.fullmatch(r"mp_(\w+)", kolom):
        return f"1 = memakai {PLATFORM.get(m.group(1), m.group(1))}"
    if m := re.fullmatch(r"(omzet|trx)_(sebelum|sesudah)(_raw|_w)?", kolom):
        var = "Omzet per bulan (Rp)" if m.group(1) == "omzet" else "Jumlah transaksi"
        jenis = {
            "_raw": "isian asli",
            "_w": "setelah nilai ekstrem dipangkas di persentil 99 (dipakai analisis)",
            None: "setelah dibersihkan",
        }[m.group(3)]
        return f"{var} {m.group(2)} marketplace, {jenis}"
    if m := re.fullmatch(r"flag_winsor_(\w+)", kolom):
        return f"1 = nilai {m.group(1)} dipangkas di persentil 99"
    return ""


def tabel_kamus(kolom_clean: list[str], kolom_feat: list[str]) -> pd.DataFrame:
    """Kamus data: kolom, sheet tempat kolom muncul, dan penjelasannya."""
    semua = list(dict.fromkeys([*kolom_clean, *kolom_feat]))
    rows = []
    for k in semua:
        sheet = " & ".join(
            s
            for s, cols in (("Data bersih", kolom_clean), ("Data siap analisis", kolom_feat))
            if k in cols
        )
        rows.append({"Kolom": k, "Ada di sheet": sheet, "Keterangan": keterangan(k)})
    return pd.DataFrame(rows)


PETUNJUK = [
    ("Data kuesioner strategi marketplace UMKM Pekanbaru", ""),
    ("", ""),
    ("Sheet", "Isi"),
    (
        "Data bersih",
        "Hasil Tahap 1: data kuesioner yang sudah dirapikan (149 responden). Nama usaha tidak disertakan.",
    ),
    (
        "Data siap analisis",
        "Hasil Tahap 2: skor strategi, target kinerja, dan profil usaha (149 responden).",
    ),
    ("Kamus data", "Penjelasan setiap kolom."),
    ("", ""),
    ("Catatan", ""),
    ("1", "Responden dikenali lewat id_responden (R001 = baris pertama di Excel kuesioner asli)."),
    ("2", "Skala jawaban kuesioner 1-5; skor strategi = rata-rata butirnya."),
    ("3", "Analisis utama memakai 116 responden dengan sampel_utama = 1."),
    (
        "4",
        "Kolom berakhiran _raw = isian asli; _w = nilai yang dipakai analisis (ekstrem dipangkas).",
    ),
    (
        "5",
        "Kolom flag_ = penanda (1 = ya, 0 = tidak); data yang ditandai tidak diubah kecuali disebut.",
    ),
    ("6", "Hasil analisis menggambarkan keterkaitan, bukan sebab-akibat."),
]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="DCE6F1")


def _rapikan_sheet(ws, df: pd.DataFrame) -> None:
    """Header tebal, baris pertama dibekukan, filter, lebar kolom, dan format angka."""
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row == 1)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        huruf = get_column_letter(i)
        sampel = df[col].astype(str).head(200).str.len()
        ws.column_dimensions[huruf].width = min(
            45, max(10, len(col) + 2, int(sampel.max() or 0) + 2)
        )
        if col.endswith("_raw") or not pd.api.types.is_numeric_dtype(df[col]):
            continue
        if re.match(r"(omzet|trx)_", col):
            fmt = "#,##0"
        elif pd.api.types.is_integer_dtype(df[col]) or col.startswith(("flag_", "mp_")):
            fmt = "0"
        else:
            fmt = "0.00"
        for cell in ws[huruf][1:]:
            cell.number_format = fmt


def ekspor_xlsx(path: Path | None = None) -> Path:
    """Tulis data_umkm.xlsx dari data_clean.csv dan data_features.csv; kembalikan path-nya."""
    proc = config.path("processed")
    clean = pd.read_csv(proc / "data_clean.csv").drop(columns=KOLOM_PRIBADI, errors="ignore")
    feat = pd.read_csv(proc / "data_features.csv").drop(columns=KOLOM_PRIBADI, errors="ignore")
    kamus = tabel_kamus(list(clean.columns), list(feat.columns))
    path = path or proc / NAMA_BERKAS

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(PETUNJUK).to_excel(xw, sheet_name="Petunjuk", index=False, header=False)
        clean.to_excel(xw, sheet_name="Data bersih", index=False)
        feat.to_excel(xw, sheet_name="Data siap analisis", index=False)
        kamus.to_excel(xw, sheet_name="Kamus data", index=False)

        ws = xw.sheets["Petunjuk"]
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=FONT, size=10, bold=c.row in (1, 3, 8))
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(name=FONT, size=13, bold=True)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 95
        for nama, df in (("Data bersih", clean), ("Data siap analisis", feat)):
            _rapikan_sheet(xw.sheets[nama], df)
        wk = xw.sheets["Kamus data"]
        _rapikan_sheet(wk, kamus)
        wk.freeze_panes = "A2"
        wk.column_dimensions["C"].width = 95
    return path


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Ekspor data bersih dan data fitur ke Excel.")
    parser.add_argument("--out", type=Path, default=None, help="path berkas .xlsx keluaran")
    args = parser.parse_args(argv)
    print(f"Ditulis: {ekspor_xlsx(args.out)}")


if __name__ == "__main__":
    main()
